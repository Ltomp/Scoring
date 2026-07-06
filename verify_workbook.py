#!/usr/bin/env python3
"""End-to-end verification of Scoring_Template_V2.xlsx.

Fills a synthetic trip (12 players, 3 of 10 rounds, an absentee, card
penalties, and deliberately engineered countback ties) into a copy of the
template, recalculates it with headless LibreOffice, then compares every
computed value against an independent Python implementation of the
competition rules. Also checks the blank template recalculates without
any spreadsheet errors.

Usage:  python3 verify_workbook.py [template.xlsx]
"""

import math
import random
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

import generate_workbook as G

HOLES = 18


# ------------------------------------------------------------- model
def round_half(x):
    """Excel ROUND: half away from zero."""
    return math.copysign(math.floor(abs(x) + 0.5), x)


def strokes(si, hc):
    return (si <= hc) + (si + 18 <= hc) + (si + 36 <= hc)


def hole_pts(score, par, si, daily_hc):
    if not score:
        return 0
    return max(0, par + 2 - score + strokes(si, daily_hc))


def adjustment(pos, n):
    if n <= 1:
        return 0
    return 2 * (2 * (pos - 1) / (n - 1) - 1)


class Model:
    """Independent implementation of the trip competition rules."""

    def __init__(self, names, start_hcaps, courses, scores, penalties):
        self.names = names
        self.n = len(names)
        self.courses = courses          # {round: (pars, sis)}
        self.scores = scores            # {round: {player_idx: [18 scores]}}
        self.penalties = penalties      # {round: {player_idx: pen}}
        self.rounds = sorted(courses)
        self.hc = {}                    # (r, p) -> running handicap into round r
        self.daily = {}                 # (r, p) -> rounded daily handicap
        self.pts = {}                   # (r, p) -> [18 hole points]
        self.raw = {}
        self.net = {}
        self.played = {}
        self.key = {}
        self.pos = {}
        self.avg = {}                   # r -> awarded average
        for p in range(self.n):
            self.hc[(1, p)] = start_hcaps[p]
        self._run()

    def _run(self):
        max_r = max(self.rounds)
        for r in range(1, max_r + 1):
            pars, sis = self.courses[r]
            for p in range(self.n):
                self.daily[(r, p)] = int(round_half(self.hc[(r, p)]))
                sc = self.scores.get(r, {}).get(p, [0] * HOLES)
                self.pts[(r, p)] = [hole_pts(sc[h], pars[h], sis[h], self.daily[(r, p)])
                                    for h in range(HOLES)]
                self.raw[(r, p)] = sum(self.pts[(r, p)])
                self.played[(r, p)] = any(sc)
            players_in = [p for p in range(self.n) if self.played[(r, p)]]
            self.avg[r] = (int(round_half(sum(self.raw[(r, p)] for p in players_in) / len(players_in)))
                           if players_in else 0)
            for p in range(self.n):
                pen = self.penalties.get(r, {}).get(p, 0)
                self.net[(r, p)] = (max(0, self.raw[(r, p)] - pen) if self.played[(r, p)]
                                    else self.avg[r])
                pts = self.pts[(r, p)]
                b9, l6, l3 = sum(pts[9:]), sum(pts[12:]), sum(pts[15:])
                h = pts  # h[17] = hole 18 ... h[12] = hole 13
                self.key[(r, p)] = (self.net[(r, p)] * 10**12 + b9 * 10**10 + l6 * 10**8
                                    + l3 * 10**6 + h[17] * 10**5 + h[16] * 10**4
                                    + h[15] * 10**3 + h[14] * 10**2 + h[13] * 10 + h[12])
            for p in range(self.n):
                self.pos[(r, p)] = 1 + sum(self.key[(r, q)] > self.key[(r, p)]
                                           for q in range(self.n))
            active = bool(players_in)
            for p in range(self.n):
                nxt = self.hc[(r, p)] + (adjustment(self.pos[(r, p)], self.n) if active else 0)
                self.hc[(r + 1, p)] = nxt

    def total(self, p):
        return sum(self.net[(r, p)] for r in self.rounds)

    def overall_pos(self, p):
        last = max(self.rounds)
        me = (self.total(p), self.key[(last, p)])
        return 1 + sum((self.total(q), self.key[(last, q)]) > me for q in range(self.n))


# ------------------------------------------------------------- scenario
def build_scenario():
    rng = random.Random(42)
    names = [f"Player {i:02d}" for i in range(1, 13)]
    start = [6, 8, 12.5, 14, 15, 16, 18, 21, 25, 27, 33, 38]

    def si_perm(seed):
        si = list(range(1, 19))
        random.Random(seed).shuffle(si)
        return si

    courses = {
        1: ([4, 4, 3, 5, 4, 3, 4, 5, 4, 4, 3, 5, 4, 4, 4, 3, 5, 4], si_perm(1)),
        2: ([5, 4, 4, 3, 4, 4, 5, 3, 4, 3, 4, 4, 5, 4, 3, 4, 4, 5], si_perm(2)),
        3: ([4, 3, 4, 4, 5, 4, 3, 4, 5, 4, 4, 3, 4, 5, 4, 4, 3, 4], si_perm(3)),
    }

    def rand_card(par):
        card = []
        for h in range(HOLES):
            if rng.random() < 0.06:
                card.append(0)                       # wiped hole
            else:
                card.append(max(1, par[h] + rng.randint(-1, 4)))
        if not any(card):
            card[0] = par[0]
        return card

    scores = {1: {}, 2: {}, 3: {}}
    for p in range(12):
        scores[1][p] = rand_card(courses[1][0])
    for p in range(12):
        if p != 4:                                   # Player 05 absent in round 2
            scores[2][p] = rand_card(courses[2][0])

    penalties = {1: {1: 1, 6: 2}, 3: {2: 2}}         # round 1: P02 -1, P07 -2; round 3: P03 -2

    # Rounds 1-2 fix the handicaps going into round 3; craft round-3 cards
    # from target points so ties land exactly where we want them.
    pre = Model(names, start, {r: courses[r] for r in (1, 2)},
                {r: scores[r] for r in (1, 2)}, penalties)
    hc3 = {p: int(round_half(pre.hc[(3, p)])) for p in range(12)}
    pars3, sis3 = courses[3]

    def card_from_pts(target, hc):
        card = []
        for h in range(HOLES):
            s = pars3[h] + strokes(sis3[h], hc) + 2 - target[h]
            assert 1 <= s <= 15, (h, s)
            card.append(s)
        return card

    # P11 vs P12: equal total/b9/l6/l3/h18 — decided at hole 17
    p11 = [3, 2, 2, 2, 2, 2, 2, 2, 2, 2, 2, 2, 2, 2, 2, 3, 2, 2]   # h16=3, h17=2
    p12 = [3, 2, 2, 2, 2, 2, 2, 2, 2, 2, 2, 2, 2, 2, 2, 2, 3, 2]   # h16=2, h17=3
    # P01 vs P02: equal total — decided on back 9
    p01 = [1, 1, 2, 1, 2, 1, 1, 2, 2, 3, 2, 2, 2, 3, 2, 2, 3, 1]   # front 13 back 20
    p02 = [3, 2, 2, 2, 3, 2, 2, 2, 2, 2, 1, 2, 1, 2, 1, 2, 1, 1]   # front 20 back 13
    for p in range(12):
        scores[3][p] = rand_card(pars3)
    scores[3][10] = card_from_pts(p11, hc3[10])
    scores[3][11] = card_from_pts(p12, hc3[11])
    scores[3][0] = card_from_pts(p01, hc3[0])
    scores[3][1] = card_from_pts(p02, hc3[1])

    model = Model(names, start, courses, scores, penalties)
    # sanity: the engineered ties are real ties on net points
    assert model.net[(3, 10)] == model.net[(3, 11)]
    assert model.pos[(3, 11)] == model.pos[(3, 10)] - 1          # P12 wins at hole 17
    assert model.net[(3, 0)] == model.net[(3, 1)]
    assert model.pos[(3, 0)] < model.pos[(3, 1)]                 # P01 wins on back 9
    return names, start, courses, scores, penalties, model


# ------------------------------------------------------------- workbook I/O
def fill_workbook(template, out_path, names, start, courses, scores, penalties):
    wb = openpyxl.load_workbook(template)
    su = wb["Setup"]
    su["B1"] = "Verification Trip"
    su["B2"] = 2026
    for p, (nm, h) in enumerate(zip(names, start)):
        su.cell(row=G.SU_PLAYER_ROW0 + p, column=2).value = nm
        su.cell(row=G.SU_PLAYER_ROW0 + p, column=5).value = h
    for r, (pars, sis) in courses.items():
        pc = G.su_par_col(r)
        su.cell(row=G.SU_COURSE_NAME_ROW, column=pc).value = f"Course {r}"
        for h in range(HOLES):
            su.cell(row=G.su_hole_row(h + 1), column=pc).value = pars[h]
            su.cell(row=G.su_hole_row(h + 1), column=pc + 1).value = sis[h]
    for r in courses:
        ws = wb[f"Round {r}"]
        for p, card in scores.get(r, {}).items():
            col = G.score_col(p + 1)
            for h in range(HOLES):
                if card[h]:
                    ws.cell(row=G.hole_row(h + 1), column=col).value = card[h]
        for p, pen in penalties.get(r, {}).items():
            ws.cell(row=G.RS_PEN_ROW, column=G.score_col(p + 1)).value = pen
    wb.calculation.fullCalcOnLoad = True
    wb.save(out_path)


def recalc(path, workdir):
    """Recalculate a workbook with headless LibreOffice; returns path to result."""
    profile = Path(workdir) / "lo-profile"
    out = Path(workdir) / "out"
    out.mkdir(exist_ok=True)
    subprocess.run(
        ["soffice", "--headless", "--norestore",
         f"-env:UserInstallation=file://{profile}",
         "--convert-to", "xlsx", "--outdir", str(out), str(path)],
        check=True, capture_output=True, timeout=300)
    result = out / Path(path).name
    if not result.exists():
        raise RuntimeError(f"LibreOffice produced no output for {path}")
    return result


ERRORS = ("#REF!", "#N/A", "#VALUE!", "#DIV/0!", "#NAME?", "#NUM!", "#NULL!")


def scan_errors(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    bad = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value in ERRORS:
                    bad.append(f"{ws.title}!{c.coordinate}: {c.value}")
    return bad


# ------------------------------------------------------------- assertions
class Checker:
    def __init__(self):
        self.failures = []
        self.count = 0

    def eq(self, what, got, want, tol=0.0):
        self.count += 1
        ok = (got == want) if tol == 0 else (
            isinstance(got, (int, float)) and abs(got - want) <= tol)
        if not ok:
            self.failures.append(f"{what}: workbook={got!r} expected={want!r}")

    def report(self):
        print(f"{self.count} checks, {len(self.failures)} failures")
        for f in self.failures:
            print("  FAIL", f)
        return not self.failures


def verify_values(path, names, courses, scores, model):
    wb = openpyxl.load_workbook(path, data_only=True)
    ck = Checker()
    n = len(names)

    for r in courses:
        ws = wb[f"Round {r}"]
        for p in range(n):
            col = G.score_col(p + 1)
            tag = f"R{r} {names[p]}"
            ck.eq(f"{tag} daily hcap", ws.cell(row=G.RS_HCAP_ROW, column=col).value,
                  model.daily[(r, p)])
            for h in range(HOLES):
                ck.eq(f"{tag} pts hole {h+1}",
                      ws.cell(row=G.hole_row(h + 1), column=col + 1).value,
                      model.pts[(r, p)][h])
            ck.eq(f"{tag} raw total", ws.cell(row=G.RS_TOTAL_ROW, column=col + 1).value,
                  model.raw[(r, p)])
            ck.eq(f"{tag} net", ws.cell(row=G.RS_RNDPTS_ROW, column=col).value,
                  model.net[(r, p)])
            ck.eq(f"{tag} position", ws.cell(row=G.RS_POS_ROW, column=col).value,
                  model.pos[(r, p)])
        ck.eq(f"R{r} awarded avg", ws["T31"].value, model.avg[r])

    hc = wb["Handicapping"]
    ck.eq("players on trip", hc["B1"].value, n)
    for p in range(n):
        row = G.HC_PLAYER_ROW0 + p
        for r in range(1, len(courses) + 1):
            ck.eq(f"HC {names[p]} into rd {r}",
                  hc.cell(row=row, column=G.HC_RD1_COL + r - 1).value,
                  model.hc[(r, p)], tol=1e-9)
    for pos in range(1, n + 1):
        ck.eq(f"adjustment pos {pos}", hc.cell(row=G.HC_ADJ_ROW0 + pos - 1, column=2).value,
              adjustment(pos, n), tol=1e-9)

    lb = wb["Leaderboard"]
    for p in range(n):
        row = G.LB_PLAYER_ROW0 + p
        ck.eq(f"LB {names[p]} total", lb.cell(row=row, column=3).value, model.total(p))
        ck.eq(f"LB {names[p]} overall pos", lb.cell(row=row, column=1).value,
              model.overall_pos(p))
        for r in courses:
            ck.eq(f"LB {names[p]} rd {r} pts", lb.cell(row=row, column=2 + 2 * r).value,
                  model.net[(r, p)])
            ck.eq(f"LB {names[p]} rd {r} pos", lb.cell(row=row, column=3 + 2 * r).value,
                  model.pos[(r, p)])
        # unused rounds stay blank
        ck.eq(f"LB {names[p]} rd 5 blank", lb.cell(row=row, column=12).value, None)
    # blank roster rows stay blank
    ck.eq("LB row 34 blank", lb.cell(row=G.LB_PLAYER_ROW0 + 31, column=3).value, None)
    return ck


def main():
    template = sys.argv[1] if len(sys.argv) > 1 else "Scoring_Template_V2.xlsx"
    names, start, courses, scores, penalties, model = build_scenario()

    with tempfile.TemporaryDirectory() as td:
        # 1. blank template: no errors anywhere
        blank = recalc(template, td)
        bad = scan_errors(blank)
        if bad:
            print("blank template has errors:")
            for b in bad[:40]:
                print("  ", b)
            sys.exit(1)
        print("blank template: no spreadsheet errors")

        # 2. filled scenario matches the independent model
        filled = Path(td) / "filled.xlsx"
        fill_workbook(template, filled, names, start, courses, scores, penalties)
        result = recalc(filled, td)
        bad = [b for b in scan_errors(result)]
        if bad:
            print("filled workbook has errors:")
            for b in bad[:40]:
                print("  ", b)
            sys.exit(1)
        ck = verify_values(result, names, courses, scores, model)
        if not ck.report():
            shutil.copy(result, "verify_failed.xlsx")
            print("recalculated workbook saved to verify_failed.xlsx for inspection")
            sys.exit(1)
    print("ALL CHECKS PASSED")


if __name__ == "__main__":
    main()
