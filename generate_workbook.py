#!/usr/bin/env python3
"""Generate Scoring_Template_V2.xlsx — golf trip scoring workbook.

Builds a locked-down, formula-driven workbook for running a golf trip
Stableford competition with a custom position-based handicapping system:

  * Setup sheet        — trip inputs: players, starting handicaps, courses (par/SI)
  * Round 1..10 sheets — daily score entry, Stableford points, card penalties,
                         Australian-countback positions
  * Handicapping       — running trip handicap per round + auto-scaling
                         position->adjustment table
  * Leaderboard        — per-round points/positions and overall standings
  * Instructions       — how to run a trip

Only classic spreadsheet functions are used (INDEX/CHOOSE/SUMPRODUCT/COUNTIF...)
so the workbook behaves identically in Excel desktop, Excel on the web,
LibreOffice and Google Sheets. All calculation cells are locked; only input
cells (highlighted yellow) are editable. Regenerate with:  python3 generate_workbook.py
"""

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# ---------------------------------------------------------------- constants
MAX_PLAYERS = 32
MAX_ROUNDS = 10
HOLES = 18

# Setup sheet geometry
SU_PLAYER_HDR_ROW = 4                 # header row of players table
SU_PLAYER_ROW0 = 5                    # first player row (players 1..32 -> rows 5..36)
SU_COURSE_TITLE_ROW = 2               # courses block sits beside the players table
SU_ROUND_HDR_ROW = 3                  # "Round r" merged header
SU_COURSE_NAME_ROW = 4                # course name input (merged pair)
SU_SUBHDR_ROW = 5                     # Par / SI
SU_HOLE_ROW0 = 6                      # hole 1 row (holes 1..18 -> rows 6..23)
SU_PAR_TOTAL_ROW = 24
SU_SI_CHECK_ROW = 25
SU_LBL_COL = 6                        # column F: hole numbers / labels for course block

# Round sheet geometry
RS_TITLE_ROW = 1
RS_NAME_ROW = 2                       # player names (merged score+pts pair)
RS_HCAP_ROW = 3                       # daily handicap (merged pair)
RS_HDR_ROW = 4                        # Hole / Par / SI / Score / Pts
RS_H1_ROW0 = 5                        # holes 1..9  -> rows 5..13
RS_OUT_ROW = 14
RS_H10_ROW0 = 15                      # holes 10..18 -> rows 15..23
RS_IN_ROW = 24
RS_TOTAL_ROW = 25
RS_PEN_ROW = 26                       # card penalty input (score column)
RS_RNDPTS_ROW = 27                    # net round points (merged pair)
RS_POS_ROW = 28                       # daily position (merged pair)
RS_CB_ROW = 29                        # countback marker (merged pair)
RS_HELP_HDR_ROW = 31                  # hidden helper block header
RS_HELP_ROW0 = 32                     # helper rows: players 1..32 -> rows 32..63
RS_HELP_LAST = RS_HELP_ROW0 + MAX_PLAYERS - 1        # 63
# helper columns: A idx, B name, C played, D raw, E pen, F net, G b9, H l6,
#                 I l3, J h18, K h17, L h16, M h15, N h14, O h13, P key, Q pos
# S31 = count of players who played, T31 = awarded average for absentees

# Handicapping sheet geometry
HC_PLAYER_ROW0 = 4                    # players 1..32 -> rows 4..35
HC_START_COL = 3                      # C = starting handicap
HC_RD1_COL = 4                        # D..M = handicap going into rounds 1..10
HC_END_COL = HC_RD1_COL + MAX_ROUNDS  # N = end handicap
HC_TOTAL_COL = HC_END_COL + 1         # O = total +/-
HC_ADJ_ROW0 = 40                      # adjustment table positions 1..32 -> rows 40..71

# Leaderboard geometry
LB_HDR_ROW = 2
LB_PLAYER_ROW0 = 3                    # players 1..32 -> rows 3..34
LB_KEY_COL = 25                       # hidden col Y: countback key of last active round

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HDR_FILL = PatternFill("solid", fgColor="1F5C33")     # dark green
SUB_FILL = PatternFill("solid", fgColor="DCE9DF")     # pale green
SUM_FILL = PatternFill("solid", fgColor="EDEDED")     # grey summary rows
INPUT_FILL = PatternFill("solid", fgColor="FFF6D6")   # pale yellow = editable
HDR_FONT = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")
UNLOCKED = Protection(locked=False)


def score_col(p):
    """Score-entry column index for player p (1-based)."""
    return 4 + 2 * (p - 1)


def su_par_col(r):
    # course block sits right of the players table so column widths stay independent
    return 7 + 2 * (r - 1)


def hole_row(h):
    """Round-sheet row for hole h (1..18)."""
    return RS_H1_ROW0 + h - 1 if h <= 9 else RS_H10_ROW0 + h - 10


def su_hole_row(h):
    return SU_HOLE_ROW0 + h - 1


def style(cell, fill=None, font=None, align=CENTER, border=BORDER, fmt=None, unlocked=False):
    if fill:
        cell.fill = fill
    if font:
        cell.font = font
    if align:
        cell.alignment = align
    if border:
        cell.border = border
    if fmt:
        cell.number_format = fmt
    if unlocked:
        cell.protection = UNLOCKED
    return cell


def protect(ws):
    ws.protection.sheet = True
    ws.protection.enable()


# =========================================================== Setup sheet
def build_setup(wb):
    ws = wb.create_sheet("Setup")
    ws.sheet_properties.tabColor = "1F5C33"

    style(ws["A1"], font=BOLD, align=LEFT, border=None).value = "Trip Name:"
    style(ws["B1"], fill=INPUT_FILL, align=LEFT, unlocked=True)
    ws.merge_cells("B1:E1")
    style(ws["A2"], font=BOLD, align=LEFT, border=None).value = "Year:"
    style(ws["B2"], fill=INPUT_FILL, align=LEFT, unlocked=True)

    # ---- players table
    hdrs = ["No.", "Player Name", "Golflink No.", "Nickname", "Starting H'cap"]
    for i, h in enumerate(hdrs, start=1):
        style(ws.cell(row=SU_PLAYER_HDR_ROW, column=i), fill=HDR_FILL, font=HDR_FONT).value = h
    for p in range(1, MAX_PLAYERS + 1):
        r = SU_PLAYER_ROW0 + p - 1
        style(ws.cell(row=r, column=1), fill=SUB_FILL).value = p
        style(ws.cell(row=r, column=2), fill=INPUT_FILL, align=LEFT, unlocked=True)
        style(ws.cell(row=r, column=3), fill=INPUT_FILL, unlocked=True, fmt="0")
        style(ws.cell(row=r, column=4), fill=INPUT_FILL, align=LEFT, unlocked=True)
        style(ws.cell(row=r, column=5), fill=INPUT_FILL, unlocked=True, fmt="0.0")

    dv_hcap = DataValidation(
        type="decimal", operator="between", formula1="-5", formula2="54",
        allow_blank=True, showErrorMessage=True,
        errorTitle="Invalid handicap", error="Starting handicap must be between -5 and 54.")
    ws.add_data_validation(dv_hcap)
    dv_hcap.add(f"E{SU_PLAYER_ROW0}:E{SU_PLAYER_ROW0 + MAX_PLAYERS - 1}")

    # ---- courses block
    style(ws.cell(row=SU_COURSE_TITLE_ROW, column=SU_LBL_COL + 1), font=BOLD, align=LEFT, border=None).value = \
        "Courses — enter course name, then Par and Stroke Index for each hole"

    dv_par = DataValidation(
        type="whole", operator="between", formula1="3", formula2="6",
        allow_blank=True, showErrorMessage=True,
        errorTitle="Invalid par", error="Par must be a whole number between 3 and 6.")
    dv_si = DataValidation(
        type="whole", operator="between", formula1="1", formula2="18",
        allow_blank=True, showErrorMessage=True,
        errorTitle="Invalid stroke index", error="Stroke index must be a whole number 1-18.")
    ws.add_data_validation(dv_par)
    ws.add_data_validation(dv_si)

    style(ws.cell(row=SU_SUBHDR_ROW, column=SU_LBL_COL), fill=HDR_FILL, font=HDR_FONT).value = "Hole"
    style(ws.cell(row=SU_COURSE_NAME_ROW, column=SU_LBL_COL), font=BOLD, align=LEFT, border=None).value = "Course:"
    style(ws.cell(row=SU_PAR_TOTAL_ROW, column=SU_LBL_COL), font=BOLD, align=LEFT, border=None).value = "Par total"
    style(ws.cell(row=SU_SI_CHECK_ROW, column=SU_LBL_COL), font=BOLD, align=LEFT, border=None).value = "SI check"
    for h in range(1, HOLES + 1):
        style(ws.cell(row=su_hole_row(h), column=SU_LBL_COL), fill=SUB_FILL).value = h

    for r in range(1, MAX_ROUNDS + 1):
        pc, sc = su_par_col(r), su_par_col(r) + 1
        pl, sl = get_column_letter(pc), get_column_letter(sc)
        c = style(ws.cell(row=SU_ROUND_HDR_ROW, column=pc), fill=HDR_FILL, font=HDR_FONT)
        c.value = f"Round {r}"
        style(ws.cell(row=SU_ROUND_HDR_ROW, column=sc), fill=HDR_FILL)
        ws.merge_cells(start_row=SU_ROUND_HDR_ROW, start_column=pc,
                       end_row=SU_ROUND_HDR_ROW, end_column=sc)
        style(ws.cell(row=SU_COURSE_NAME_ROW, column=pc), fill=INPUT_FILL, align=LEFT, unlocked=True)
        style(ws.cell(row=SU_COURSE_NAME_ROW, column=sc), fill=INPUT_FILL, unlocked=True)
        ws.merge_cells(start_row=SU_COURSE_NAME_ROW, start_column=pc,
                       end_row=SU_COURSE_NAME_ROW, end_column=sc)
        style(ws.cell(row=SU_SUBHDR_ROW, column=pc), fill=SUB_FILL, font=BOLD).value = "Par"
        style(ws.cell(row=SU_SUBHDR_ROW, column=sc), fill=SUB_FILL, font=BOLD).value = "SI"
        for h in range(1, HOLES + 1):
            row = su_hole_row(h)
            style(ws.cell(row=row, column=pc), fill=INPUT_FILL, unlocked=True, fmt="0")
            style(ws.cell(row=row, column=sc), fill=INPUT_FILL, unlocked=True, fmt="0")
        dv_par.add(f"{pl}{SU_HOLE_ROW0}:{pl}{SU_HOLE_ROW0 + HOLES - 1}")
        dv_si.add(f"{sl}{SU_HOLE_ROW0}:{sl}{SU_HOLE_ROW0 + HOLES - 1}")

        par_rng = f"{pl}{SU_HOLE_ROW0}:{pl}{SU_HOLE_ROW0 + HOLES - 1}"
        si_rng = f"{sl}{SU_HOLE_ROW0}:{sl}{SU_HOLE_ROW0 + HOLES - 1}"
        style(ws.cell(row=SU_PAR_TOTAL_ROW, column=pc), fill=SUM_FILL, font=BOLD).value = \
            f'=IF(COUNT({par_rng})=0,"",SUM({par_rng}))'
        chk = style(ws.cell(row=SU_SI_CHECK_ROW, column=pc), fill=SUM_FILL, font=BOLD)
        chk.value = (f'=IF(COUNT({si_rng})=0,"",IF(AND(COUNT({si_rng})=18,SUM({si_rng})=171,'
                     f'SUMPRODUCT(--(COUNTIF({si_rng},{si_rng})>1))=0),"OK","CHECK"))')
        ws.merge_cells(start_row=SU_SI_CHECK_ROW, start_column=pc,
                       end_row=SU_SI_CHECK_ROW, end_column=sc)
        style(ws.cell(row=SU_SI_CHECK_ROW, column=sc), fill=SUM_FILL)

        # highlight duplicate SI entries within the round
        ws.conditional_formatting.add(
            si_rng,
            FormulaRule(formula=[f'AND({sl}{SU_HOLE_ROW0}<>"",COUNTIF({sl}${SU_HOLE_ROW0}:{sl}${SU_HOLE_ROW0 + HOLES - 1},{sl}{SU_HOLE_ROW0})>1)'],
                        fill=PatternFill("solid", fgColor="F4CCCC"), stopIfTrue=False))
    # flag failed SI check
    first = get_column_letter(su_par_col(1))
    ws.conditional_formatting.add(
        f"{first}{SU_SI_CHECK_ROW}:{get_column_letter(su_par_col(MAX_ROUNDS)+1)}{SU_SI_CHECK_ROW}",
        FormulaRule(formula=[f'{first}{SU_SI_CHECK_ROW}="CHECK"'],
                    fill=PatternFill("solid", fgColor="F4CCCC"), stopIfTrue=False))

    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 13
    ws.column_dimensions["D"].width = 13
    ws.column_dimensions["E"].width = 14
    ws.column_dimensions[get_column_letter(SU_LBL_COL)].width = 9
    for r in range(1, MAX_ROUNDS + 1):
        ws.column_dimensions[get_column_letter(su_par_col(r))].width = 7
        ws.column_dimensions[get_column_letter(su_par_col(r) + 1)].width = 7
    protect(ws)
    return ws


# =========================================================== Round sheets
def build_round(wb, r):
    ws = wb.create_sheet(f"Round {r}")
    ws.sheet_properties.tabColor = "2E75B6"
    pc = su_par_col(r)
    name_ref = f"Setup!{get_column_letter(pc)}${SU_COURSE_NAME_ROW}"

    style(ws["A1"], font=Font(bold=True, size=12), align=LEFT, border=None).value = f"Round {r}"
    c = style(ws["B1"], font=Font(bold=True, size=12), align=LEFT, border=None)
    c.value = f'=IF({name_ref}="","(enter course on Setup sheet)",{name_ref})'
    ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=11)

    style(ws.cell(row=RS_NAME_ROW, column=1), fill=HDR_FILL, font=HDR_FONT, align=LEFT).value = "Player"
    style(ws.cell(row=RS_HCAP_ROW, column=1), fill=SUB_FILL, font=BOLD, align=LEFT).value = "Daily H'cap"
    for col, lbl in ((1, "Hole"), (2, "Par"), (3, "SI")):
        style(ws.cell(row=RS_HDR_ROW, column=col), fill=HDR_FILL, font=HDR_FONT).value = lbl
    style(ws.cell(row=RS_NAME_ROW, column=2), fill=HDR_FILL)
    style(ws.cell(row=RS_NAME_ROW, column=3), fill=HDR_FILL)
    style(ws.cell(row=RS_HCAP_ROW, column=2), fill=SUB_FILL)
    style(ws.cell(row=RS_HCAP_ROW, column=3), fill=SUB_FILL)

    # hole rows: par / SI pulled from Setup for THIS round
    pl, sl = get_column_letter(pc), get_column_letter(pc + 1)
    for h in range(1, HOLES + 1):
        row, srow = hole_row(h), su_hole_row(h)
        style(ws.cell(row=row, column=1), fill=SUB_FILL).value = h
        style(ws.cell(row=row, column=2)).value = f'=IF(Setup!{pl}${srow}="","",Setup!{pl}${srow})'
        style(ws.cell(row=row, column=3)).value = f'=IF(Setup!{sl}${srow}="","",Setup!{sl}${srow})'

    for row, lbl in ((RS_OUT_ROW, "Out"), (RS_IN_ROW, "In"), (RS_TOTAL_ROW, "Total")):
        style(ws.cell(row=row, column=1), fill=SUM_FILL, font=BOLD).value = lbl
        b = "=SUM(B5:B13)" if row == RS_OUT_ROW else ("=SUM(B15:B23)" if row == RS_IN_ROW else "=SUM(B14,B24)")
        style(ws.cell(row=row, column=2), fill=SUM_FILL, font=BOLD).value = b.replace("SUM(B", "SUM(B")
        style(ws.cell(row=row, column=3), fill=SUM_FILL)
    ws["B14"] = "=IF(COUNT(B5:B13)=0,\"\",SUM(B5:B13))"
    ws["B24"] = "=IF(COUNT(B15:B23)=0,\"\",SUM(B15:B23))"
    ws["B25"] = "=IF(COUNT(B5:B13,B15:B23)=0,\"\",SUM(B5:B13,B15:B23))"

    style(ws.cell(row=RS_PEN_ROW, column=1), font=BOLD, align=LEFT).value = "Card penalty (0/1/2)"
    style(ws.cell(row=RS_RNDPTS_ROW, column=1), fill=SUM_FILL, font=BOLD, align=LEFT).value = "Round points"
    style(ws.cell(row=RS_POS_ROW, column=1), fill=SUM_FILL, font=BOLD, align=LEFT).value = "Daily position"
    style(ws.cell(row=RS_CB_ROW, column=1), align=LEFT).value = "Tie decided on countback"
    for col in (2, 3):
        style(ws.cell(row=RS_PEN_ROW, column=col))
        style(ws.cell(row=RS_RNDPTS_ROW, column=col), fill=SUM_FILL)
        style(ws.cell(row=RS_POS_ROW, column=col), fill=SUM_FILL)
        style(ws.cell(row=RS_CB_ROW, column=col))

    dv_score = DataValidation(
        type="whole", operator="between", formula1="0", formula2="15",
        allow_blank=True, showErrorMessage=True, errorTitle="Invalid score",
        error="Gross score must be a whole number 1-15 (0 or blank = no score / wipe).")
    dv_pen = DataValidation(
        type="list", formula1='"0,1,2"', allow_blank=True, showErrorMessage=True,
        errorTitle="Invalid penalty", error="Card penalty must be 0, 1 or 2 points.")
    ws.add_data_validation(dv_score)
    ws.add_data_validation(dv_pen)

    hd_row = RS_HELP_HDR_ROW
    for col, lbl in enumerate(["idx", "name", "played", "raw", "pen", "net", "b9", "l6",
                               "l3", "h18", "h17", "h16", "h15", "h14", "h13", "key", "pos"], start=1):
        ws.cell(row=hd_row, column=col).value = lbl
    ws[f"S{hd_row}"] = f"=SUM(C{RS_HELP_ROW0}:C{RS_HELP_LAST})"                     # players who played
    ws[f"T{hd_row}"] = (f"=IF(S{hd_row}=0,0,ROUND(SUMPRODUCT(C{RS_HELP_ROW0}:C{RS_HELP_LAST},"
                        f"D{RS_HELP_ROW0}:D{RS_HELP_LAST})/S{hd_row},0))")           # awarded average

    for p in range(1, MAX_PLAYERS + 1):
        scn = score_col(p)
        s, t = get_column_letter(scn), get_column_letter(scn + 1)
        su_row = SU_PLAYER_ROW0 + p - 1
        hr = RS_HELP_ROW0 + p - 1
        hc_row = HC_PLAYER_ROW0 + p - 1
        hc_col = get_column_letter(HC_RD1_COL + r - 1)
        scores = f"{s}5:{s}13,{s}15:{s}23"

        # header: name + daily handicap (merged over the score/pts pair)
        c = style(ws.cell(row=RS_NAME_ROW, column=scn), fill=HDR_FILL, font=HDR_FONT)
        c.value = f'=IF(Setup!$B${su_row}="","",Setup!$B${su_row})'
        style(ws.cell(row=RS_NAME_ROW, column=scn + 1), fill=HDR_FILL)
        ws.merge_cells(start_row=RS_NAME_ROW, start_column=scn, end_row=RS_NAME_ROW, end_column=scn + 1)
        c = style(ws.cell(row=RS_HCAP_ROW, column=scn), fill=SUB_FILL, font=BOLD, fmt="0")
        c.value = f'=IF({s}${RS_NAME_ROW}="","",ROUND(Handicapping!{hc_col}{hc_row},0))'
        style(ws.cell(row=RS_HCAP_ROW, column=scn + 1), fill=SUB_FILL)
        ws.merge_cells(start_row=RS_HCAP_ROW, start_column=scn, end_row=RS_HCAP_ROW, end_column=scn + 1)
        style(ws.cell(row=RS_HDR_ROW, column=scn), fill=HDR_FILL, font=HDR_FONT).value = "Score"
        style(ws.cell(row=RS_HDR_ROW, column=scn + 1), fill=HDR_FILL, font=HDR_FONT).value = "Pts"

        # hole rows: score input + points
        for h in range(1, HOLES + 1):
            row = hole_row(h)
            style(ws.cell(row=row, column=scn), fill=INPUT_FILL, unlocked=True, fmt="0")
            pts = style(ws.cell(row=row, column=scn + 1), fmt="0")
            pts.value = (f'=IF(OR(N({s}{row})=0,$B{row}="",$C{row}=""),0,'
                         f'MAX(0,$B{row}+2-{s}{row}'
                         f'+IF($C{row}<={s}${RS_HCAP_ROW},1,0)'
                         f'+IF($C{row}+18<={s}${RS_HCAP_ROW},1,0)'
                         f'+IF($C{row}+36<={s}${RS_HCAP_ROW},1,0)))')
        dv_score.add(f"{s}{RS_H1_ROW0}:{s}{RS_H1_ROW0 + 8}")
        dv_score.add(f"{s}{RS_H10_ROW0}:{s}{RS_H10_ROW0 + 8}")

        # summary rows
        style(ws.cell(row=RS_OUT_ROW, column=scn), fill=SUM_FILL, font=BOLD).value = \
            f'=IF(COUNT({s}5:{s}13)=0,"",SUM({s}5:{s}13))'
        style(ws.cell(row=RS_OUT_ROW, column=scn + 1), fill=SUM_FILL, font=BOLD).value = f"=SUM({t}5:{t}13)"
        style(ws.cell(row=RS_IN_ROW, column=scn), fill=SUM_FILL, font=BOLD).value = \
            f'=IF(COUNT({s}15:{s}23)=0,"",SUM({s}15:{s}23))'
        style(ws.cell(row=RS_IN_ROW, column=scn + 1), fill=SUM_FILL, font=BOLD).value = f"=SUM({t}15:{t}23)"
        style(ws.cell(row=RS_TOTAL_ROW, column=scn), fill=SUM_FILL, font=BOLD).value = \
            f'=IF(COUNT({scores})=0,"",SUM({scores}))'
        style(ws.cell(row=RS_TOTAL_ROW, column=scn + 1), fill=SUM_FILL, font=BOLD).value = \
            f"=SUM({t}{RS_OUT_ROW},{t}{RS_IN_ROW})"

        # penalty input
        style(ws.cell(row=RS_PEN_ROW, column=scn), fill=INPUT_FILL, unlocked=True, fmt="0")
        style(ws.cell(row=RS_PEN_ROW, column=scn + 1))
        dv_pen.add(f"{s}{RS_PEN_ROW}")

        # net round points / position / countback marker (hidden helpers do the work)
        c = style(ws.cell(row=RS_RNDPTS_ROW, column=scn), fill=SUM_FILL, font=BOLD, fmt="0")
        c.value = f'=IF(OR({s}${RS_NAME_ROW}="",$S${hd_row}=0),"",F{hr})'
        style(ws.cell(row=RS_RNDPTS_ROW, column=scn + 1), fill=SUM_FILL)
        ws.merge_cells(start_row=RS_RNDPTS_ROW, start_column=scn, end_row=RS_RNDPTS_ROW, end_column=scn + 1)
        c = style(ws.cell(row=RS_POS_ROW, column=scn), fill=SUM_FILL, font=BOLD, fmt="0")
        c.value = f'=IF(OR({s}${RS_NAME_ROW}="",$S${hd_row}=0),"",Q{hr})'
        style(ws.cell(row=RS_POS_ROW, column=scn + 1), fill=SUM_FILL)
        ws.merge_cells(start_row=RS_POS_ROW, start_column=scn, end_row=RS_POS_ROW, end_column=scn + 1)
        c = style(ws.cell(row=RS_CB_ROW, column=scn))
        c.value = (f'=IF(OR({s}${RS_NAME_ROW}="",$S${hd_row}=0),"",'
                   f'IF(SUMPRODUCT(($B${RS_HELP_ROW0}:$B${RS_HELP_LAST}<>"")'
                   f'*($F${RS_HELP_ROW0}:$F${RS_HELP_LAST}=F{hr}))>1,"cb",""))')
        style(ws.cell(row=RS_CB_ROW, column=scn + 1))
        ws.merge_cells(start_row=RS_CB_ROW, start_column=scn, end_row=RS_CB_ROW, end_column=scn + 1)

        # ---- hidden helper row (all numeric so SUMPRODUCT stays happy)
        ws.cell(row=hr, column=1).value = p
        ws.cell(row=hr, column=2).value = f'=IF({s}${RS_NAME_ROW}="","",{s}${RS_NAME_ROW})'
        ws.cell(row=hr, column=3).value = f'=IF(B{hr}="",0,IF(COUNT({scores})>0,1,0))'
        ws.cell(row=hr, column=4).value = f"=N({t}{RS_TOTAL_ROW})"
        ws.cell(row=hr, column=5).value = f"=N({s}{RS_PEN_ROW})"
        ws.cell(row=hr, column=6).value = f'=IF(B{hr}="",0,IF(C{hr}=1,MAX(0,D{hr}-E{hr}),$T${hd_row}))'
        ws.cell(row=hr, column=7).value = f"=N({t}{RS_IN_ROW})"
        ws.cell(row=hr, column=8).value = f"=SUM({t}18:{t}23)"     # holes 13-18
        ws.cell(row=hr, column=9).value = f"=SUM({t}21:{t}23)"     # holes 16-18
        for i, hrow in enumerate((23, 22, 21, 20, 19, 18)):        # h18..h13
            ws.cell(row=hr, column=10 + i).value = f"=N({t}{hrow})"
        ws.cell(row=hr, column=16).value = (
            f"=F{hr}*10^12+G{hr}*10^10+H{hr}*10^8+I{hr}*10^6"
            f"+J{hr}*100000+K{hr}*10000+L{hr}*1000+M{hr}*100+N{hr}*10+O{hr}")
        ws.cell(row=hr, column=17).value = (
            f'=IF(B{hr}="","",1+SUMPRODUCT(($B${RS_HELP_ROW0}:$B${RS_HELP_LAST}<>"")'
            f'*($P${RS_HELP_ROW0}:$P${RS_HELP_LAST}>P{hr})))')

    # hide helper block
    for row in range(RS_HELP_HDR_ROW, RS_HELP_LAST + 1):
        ws.row_dimensions[row].hidden = True

    ws.column_dimensions["A"].width = 19
    ws.column_dimensions["B"].width = 5
    ws.column_dimensions["C"].width = 5
    for p in range(1, MAX_PLAYERS + 1):
        ws.column_dimensions[get_column_letter(score_col(p))].width = 6.5
        ws.column_dimensions[get_column_letter(score_col(p) + 1)].width = 5.5
    ws.freeze_panes = "D5"
    protect(ws)
    return ws


# =========================================================== Handicapping
def build_handicapping(wb):
    ws = wb.create_sheet("Handicapping")
    ws.sheet_properties.tabColor = "7F7F7F"

    style(ws["A1"], font=BOLD, align=LEFT, border=None).value = "Players on trip:"
    style(ws["B1"], font=BOLD).value = f"=COUNTA(Setup!$B${SU_PLAYER_ROW0}:$B${SU_PLAYER_ROW0 + MAX_PLAYERS - 1})"

    hdr = ["No.", "Player", "Start H'cap"] + [f"Rd {r}" for r in range(1, MAX_ROUNDS + 1)] + \
          ["End H'cap", "Total +/-"]
    for i, h in enumerate(hdr, start=1):
        style(ws.cell(row=3, column=i), fill=HDR_FILL, font=HDR_FONT).value = h

    adj_rng = f"$B${HC_ADJ_ROW0}:$B${HC_ADJ_ROW0 + MAX_PLAYERS - 1}"
    for p in range(1, MAX_PLAYERS + 1):
        row = HC_PLAYER_ROW0 + p - 1
        su_row = SU_PLAYER_ROW0 + p - 1
        hr = RS_HELP_ROW0 + p - 1
        style(ws.cell(row=row, column=1), fill=SUB_FILL).value = p
        style(ws.cell(row=row, column=2), align=LEFT).value = f'=IF(Setup!B{su_row}="","",Setup!B{su_row})'
        style(ws.cell(row=row, column=3), fmt="0.00").value = f'=IF($B{row}="","",N(Setup!E{su_row}))'
        for r in range(1, MAX_ROUNDS + 1):
            col = HC_RD1_COL + r - 1
            cell = style(ws.cell(row=row, column=col), fmt="0.00")
            if r == 1:
                cell.value = f'=IF($B{row}="","",$C{row})'
            else:
                prev = f"{get_column_letter(col - 1)}{row}"
                prv_sheet = f"'Round {r - 1}'"
                cell.value = (f'=IF($B{row}="","",{prev}+IF({prv_sheet}!$S${RS_HELP_HDR_ROW}>0,'
                              f"N(INDEX({adj_rng},{prv_sheet}!$Q${hr})),0))")
        prev = f"{get_column_letter(HC_END_COL - 1)}{row}"
        last = f"'Round {MAX_ROUNDS}'"
        style(ws.cell(row=row, column=HC_END_COL), fmt="0.00").value = (
            f'=IF($B{row}="","",{prev}+IF({last}!$S${RS_HELP_HDR_ROW}>0,'
            f"N(INDEX({adj_rng},{last}!$Q${hr})),0))")
        style(ws.cell(row=row, column=HC_TOTAL_COL), fmt="+0.00;-0.00;0", font=BOLD).value = (
            f'=IF($B{row}="","",{get_column_letter(HC_END_COL)}{row}-$C{row})')

    # ---- position -> adjustment table (auto-scaled, overridable)
    t = style(ws.cell(row=HC_ADJ_ROW0 - 2, column=1), font=BOLD, align=LEFT, border=None)
    t.value = "Daily Handicap Adjustment (auto-scales to player count — amounts may be overridden)"
    style(ws.cell(row=HC_ADJ_ROW0 - 1, column=1), fill=HDR_FILL, font=HDR_FONT).value = "Position"
    style(ws.cell(row=HC_ADJ_ROW0 - 1, column=2), fill=HDR_FILL, font=HDR_FONT).value = "Amount"
    for p in range(1, MAX_PLAYERS + 1):
        row = HC_ADJ_ROW0 + p - 1
        style(ws.cell(row=row, column=1), fill=SUB_FILL).value = p
        amt = style(ws.cell(row=row, column=2), fill=INPUT_FILL, fmt="+0.00;-0.00;0", unlocked=True)
        amt.value = (f'=IF(OR($B$1=0,A{row}>$B$1),"",'
                     f'IF($B$1<=1,0,2*(2*(A{row}-1)/($B$1-1)-1)))')

    ws.column_dimensions["A"].width = 9
    ws.column_dimensions["B"].width = 22
    for col in range(3, HC_TOTAL_COL + 1):
        ws.column_dimensions[get_column_letter(col)].width = 9
    ws.freeze_panes = "D4"
    protect(ws)
    return ws


# =========================================================== Leaderboard
def build_leaderboard(wb):
    ws = wb.create_sheet("Leaderboard")
    ws.sheet_properties.tabColor = "BF8F00"

    style(ws["A1"], font=Font(bold=True, size=12), align=LEFT, border=None).value = "Leaderboard"
    # hidden: last active (played) round number
    terms = ",".join(f"('Round {r}'!$S${RS_HELP_HDR_ROW}>0)*{r}" for r in range(1, MAX_ROUNDS + 1))
    ws[f"{get_column_letter(LB_KEY_COL)}1"] = f"=MAX({terms})"

    style(ws.cell(row=LB_HDR_ROW, column=1), fill=HDR_FILL, font=HDR_FONT).value = "Pos"
    style(ws.cell(row=LB_HDR_ROW, column=2), fill=HDR_FILL, font=HDR_FONT).value = "Player"
    style(ws.cell(row=LB_HDR_ROW, column=3), fill=HDR_FILL, font=HDR_FONT).value = "Total"
    for r in range(1, MAX_ROUNDS + 1):
        style(ws.cell(row=LB_HDR_ROW, column=2 + 2 * r), fill=HDR_FILL, font=HDR_FONT).value = f"Rd {r}"
        style(ws.cell(row=LB_HDR_ROW, column=3 + 2 * r), fill=HDR_FILL, font=HDR_FONT).value = "Pos"

    last_row = LB_PLAYER_ROW0 + MAX_PLAYERS - 1
    key_l = get_column_letter(LB_KEY_COL)
    for p in range(1, MAX_PLAYERS + 1):
        row = LB_PLAYER_ROW0 + p - 1
        su_row = SU_PLAYER_ROW0 + p - 1
        hr = RS_HELP_ROW0 + p - 1
        style(ws.cell(row=row, column=2), align=LEFT).value = f'=IF(Setup!B{su_row}="","",Setup!B{su_row})'
        pts_cells = []
        for r in range(1, MAX_ROUNDS + 1):
            pcol, ocol = 2 + 2 * r, 3 + 2 * r
            sheet = f"'Round {r}'"
            style(ws.cell(row=row, column=pcol), fmt="0").value = (
                f'=IF($B{row}="","",IF({sheet}!$S${RS_HELP_HDR_ROW}=0,"",{sheet}!$F${hr}))')
            style(ws.cell(row=row, column=ocol), fmt="0").value = (
                f'=IF($B{row}="","",IF({sheet}!$S${RS_HELP_HDR_ROW}=0,"",{sheet}!$Q${hr}))')
            pts_cells.append(f"{get_column_letter(pcol)}{row}")
        style(ws.cell(row=row, column=3), font=BOLD, fmt="0").value = \
            f'=IF($B{row}="","",SUM({",".join(pts_cells)}))'
        choose = ",".join(f"'Round {r}'!$P${hr}" for r in range(1, MAX_ROUNDS + 1))
        ws.cell(row=row, column=LB_KEY_COL).value = \
            f'=IF(OR($B{row}="",${key_l}$1=0),0,CHOOSE(${key_l}$1,{choose}))'
        style(ws.cell(row=row, column=1), font=BOLD, fmt="0").value = (
            f'=IF($B{row}="","",1+SUMPRODUCT(($B${LB_PLAYER_ROW0}:$B${last_row}<>"")*'
            f'(($C${LB_PLAYER_ROW0}:$C${last_row}>C{row})+'
            f'(($C${LB_PLAYER_ROW0}:$C${last_row}=C{row})*'
            f'(${key_l}${LB_PLAYER_ROW0}:${key_l}${last_row}>{key_l}{row})))))')

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 8
    for r in range(1, MAX_ROUNDS + 1):
        ws.column_dimensions[get_column_letter(2 + 2 * r)].width = 6.5
        ws.column_dimensions[get_column_letter(3 + 2 * r)].width = 5.5
    ws.column_dimensions[key_l].hidden = True
    ws.freeze_panes = "D3"
    protect(ws)
    return ws


# =========================================================== Instructions
INSTRUCTIONS = [
    ("Golf Trip Scoring — how to run a trip", True),
    ("", False),
    ("BEFORE THE TRIP (Setup sheet — yellow cells are the only editable ones)", True),
    ("1.  Enter the trip name and year.", False),
    ("2.  Enter each player's name and Starting H'cap (Golflink no. and nickname are optional).", False),
    ("     Leave unused player rows completely blank — they are ignored everywhere.", False),
    ("3.  For each round, enter the course name and the Par and Stroke Index (SI) of every hole.", False),
    ("     The 'SI check' row must show OK (each of 1-18 used exactly once). Leave unused rounds blank.", False),
    ("", False),
    ("EACH DAY (Round sheets — organisers enter everything from the physical cards)", True),
    ("1.  Enter each player's gross score per hole in the yellow Score cells.", False),
    ("     Enter 0 (or leave blank) for a wiped hole — it scores no points.", False),
    ("2.  If a card was filled out incorrectly, enter a 1 or 2 point Card penalty for that player.", False),
    ("     The penalty reduces their Round points, daily position and handicap outcome,", False),
    ("     but never changes the average awarded to players who missed the round.", False),
    ("3.  A player with no scores entered is treated as absent: they receive the field's average", False),
    ("     points (rounded, before penalties) and are ranked with it in the daily comp.", False),
    ("4.  Round points, daily positions (Australian countback: back 9, last 6, last 3, then", False),
    ("     hole-by-hole from 18) and the next day's handicaps all update automatically.", False),
    ("     'cb' under a player means their tie was split by countback.", False),
    ("", False),
    ("HANDICAPPING", True),
    ("Each player's handicap moves after every round by the Amount shown for their daily", False),
    ("position (Handicapping sheet). The table auto-scales to the number of players — the", False),
    ("winner always loses 2 and last place always gains 2, with everyone between spread in", False),
    ("equal steps. You may overtype the amounts (yellow) before the trip if you want different", False),
    ("steps.", False),
    ("The Daily H'cap used on the cards is the running handicap rounded to a whole number.", False),
    ("", False),
    ("INTEGRITY", True),
    ("All white cells are locked formulas. To unlock (not recommended mid-trip):", False),
    ("Review > Unprotect Sheet (there is no password).", False),
    ("", False),
    ("OFFICE SCRIPTS (Excel on the web, optional)", True),
    ("Three helper scripts ship alongside this workbook (see office-scripts/ or the README):", False),
    ("  Tidy Trip View — hides unused player columns, rows and round sheets.", False),
    ("  New Trip Reset — clears scores, players and courses for the next trip.", False),
    ("  Lock Completed Round — locks a finished day's scores against changes.", False),
    ("Add them once via Automate > New Script, paste the code, Save Script.", False),
]


def build_instructions(wb):
    ws = wb.create_sheet("Instructions", 0)
    ws.sheet_properties.tabColor = "C00000"
    for i, (text, bold) in enumerate(INSTRUCTIONS, start=1):
        c = ws.cell(row=i, column=1)
        c.value = text
        c.font = Font(bold=bold, size=12 if i == 1 else 11)
        c.alignment = LEFT
    ws.column_dimensions["A"].width = 110
    protect(ws)
    return ws


# =========================================================== main
def main(path="Scoring_Template_V2.xlsx"):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    build_instructions(wb)
    build_setup(wb)
    build_leaderboard(wb)
    build_handicapping(wb)
    for r in range(1, MAX_ROUNDS + 1):
        build_round(wb, r)

    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    print(f"wrote {path}")


if __name__ == "__main__":
    main()
